#!/usr/bin/env python3
"""
Smart Data/References Loader
Handles directories and multiple file formats

Supports:
- Data: CSV, JSON, TXT, MD, Excel (.xlsx, .xls), DOCX
- References: PDF, TXT, MD, DOCX, CAJ (Chinese Academic Journals)
"""

import os
import sys
import json
import csv
import subprocess
from pathlib import Path

# Optional imports
try:
    import pandas as pd
    HAS_PANDAS = True
except ImportError:
    HAS_PANDAS = False

try:
    import PyPDF2
    HAS_PYPDF2 = True
except ImportError:
    HAS_PYPDF2 = False

try:
    from docx import Document
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


class SmartLoader:
    """Load and process data/references from various formats"""
    
    def __init__(self, directory, is_data_dir=True):
        self.dir = Path(directory)
        self.is_data_dir = is_data_dir
        self.content = []
        self.warnings = []
    
    def load_all(self):
        """Load all supported files from directory"""
        if not self.dir.exists():
            print(f"⚠️  Directory not found: {self.dir}")
            return ""
        
        files = list(self.dir.iterdir())
        print(f"📁 Scanning {len(files)} files in {self.dir}...")
        
        for f in files:
            if f.is_file():
                self._process_file(f)
        
        result = "\n\n".join(self.content) if self.content else ""
        
        # Print summary
        print(f"   ✅ Files loaded: {len(self.content)}")
        if self.warnings:
            print(f"   ⚠️  Warnings: {len(self.warnings)}")
            for w in self.warnings[:3]:
                print(f"      - {w}")
        
        return result
    
    def _process_file(self, filepath):
        """Process a single file based on extension"""
        ext = filepath.suffix.lower()
        
        if self.is_data_dir:
            # Loading data files
            if ext in ['.csv']:
                self._load_csv(filepath)
            elif ext in ['.json']:
                self._load_json(filepath)
            elif ext in ['.xlsx', '.xls'] and HAS_PANDAS:
                self._load_excel(filepath)
            elif ext == '.docx' and HAS_DOCX:
                self._load_docx_data(filepath)
            elif ext in ['.txt', '.md', '.data']:
                self._load_text(filepath)
            elif ext == '.caj':
                self._handle_caj(filepath, is_data=True)
            else:
                self.warnings.append(f"Unsupported format: {filepath.name}")
        else:
            # Loading reference files
            if ext == '.pdf' and HAS_PYPDF2:
                self._load_pdf(filepath)
            elif ext == '.docx' and HAS_DOCX:
                self._load_docx_refs(filepath)
            elif ext == '.caj':
                self._handle_caj(filepath, is_data=False)
            elif ext in ['.txt', '.md', '.bib', '.refs']:
                self._load_text(filepath)
            else:
                self.warnings.append(f"Unsupported format: {filepath.name}")
    
    def _load_csv(self, filepath):
        """Load CSV file"""
        try:
            with open(filepath, 'r', encoding='utf-8') as f:
                reader = csv.reader(f)
                rows = list(reader)
                
            if rows:
                content = f"### Data from {filepath.name}\n\n"
                content += "| " + " | ".join(rows[0]) + " |\n"
                content += "|" + "|".join(["---" for _ in rows[0]]) + "|\n"
                for row in rows[1:20]:
                    content += "| " + " | ".join(row) + " |\n"
                if len(rows) > 20:
                    content += f"\n*({len(rows) - 20} more rows)*\n"
                
                self.content.append(content)
                print(f"   📊 Loaded CSV: {filepath.name}")
        except Exception as e:
            self.warnings.append(f"CSV error {filepath.name}: {e}")
    
    def _load_json(self, filepath):
        """Load JSON file"""
        try:
            with open(filepath, 'r', encoding='utf-8') as f:
                data = json.load(f)
            
            content = f"### Data from {filepath.name}\n\n```json\n"
            content += json.dumps(data, indent=2)[:5000]
            content += "\n```\n"
            
            self.content.append(content)
            print(f"   📊 Loaded JSON: {filepath.name}")
        except Exception as e:
            self.warnings.append(f"JSON error {filepath.name}: {e}")
    
    def _load_excel(self, filepath):
        """Load Excel file"""
        try:
            df = pd.read_excel(filepath)
            content = f"### Data from {filepath.name}\n\n"
            content += df.to_markdown(index=False)[:5000]
            self.content.append(content)
            print(f"   📊 Loaded Excel: {filepath.name}")
        except Exception as e:
            self.warnings.append(f"Excel error {filepath.name}: {e}")
    
    def _load_docx_data(self, filepath):
        """Load DOCX as data source"""
        if not HAS_DOCX:
            self.warnings.append("python-docx not installed")
            return
        
        try:
            doc = Document(filepath)
            text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
            
            # Try to extract tables
            tables_data = []
            for table in doc.tables[:3]:  # First 3 tables
                table_data = []
                for row in table.rows:
                    row_data = [cell.text for cell in row.cells]
                    table_data.append(row_data)
                tables_data.append(table_data)
            
            content = f"### Data from {filepath.name}\n\n"
            
            # Add tables if found
            if tables_data:
                for i, table in enumerate(tables_data):
                    if table:
                        content += f"**Table {i+1}:**\n\n"
                        content += "| " + " | ".join(table[0]) + " |\n"
                        content += "|" + "|".join(["---" for _ in table[0]]) + "|\n"
                        for row in table[1:15]:
                            content += "| " + " | ".join(row) + " |\n"
                        content += "\n"
            
            # Add text content
            if text:
                content += f"**Notes:**\n\n{text[:3000]}\n"
            
            self.content.append(content)
            print(f"   📄 Loaded DOCX (data): {filepath.name}")
        except Exception as e:
            self.warnings.append(f"DOCX error {filepath.name}: {e}")
    
    def _load_docx_refs(self, filepath):
        """Load DOCX as references"""
        if not HAS_DOCX:
            self.warnings.append("python-docx not installed")
            return
        
        try:
            doc = Document(filepath)
            text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
            
            content = f"### Reference from {filepath.name}\n\n{text[:8000]}"
            self.content.append(content)
            print(f"   📄 Loaded DOCX (refs): {filepath.name}")
        except Exception as e:
            self.warnings.append(f"DOCX error {filepath.name}: {e}")
    
    def _load_pdf(self, filepath):
        """Extract text from PDF"""
        if not HAS_PYPDF2:
            self.warnings.append("PyPDF2 not installed")
            return
        
        try:
            text = ""
            with open(filepath, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages[:5]:
                    text += page.extract_text() + "\n"
            
            content = f"### Reference from {filepath.name}\n\n{text[:8000]}"
            self.content.append(content)
            print(f"   📄 Loaded PDF: {filepath.name}")
        except Exception as e:
            self.warnings.append(f"PDF error {filepath.name}: {e}")
    
    def _handle_caj(self, filepath, is_data=False):
        """Handle CAJ (Chinese Academic Journal) files"""
        label = "Data" if is_data else "Reference"
        
        # Check for conversion tools
        caj2pdf = subprocess.run(['which', 'caj2pdf'], capture_output=True).returncode == 0
        
        if caj2pdf:
            try:
                # Try to convert CAJ to PDF
                temp_pdf = filepath.with_suffix('.pdf')
                subprocess.run(['caj2pdf', '-i', str(filepath), '-o', str(temp_pdf)], 
                              check=True, capture_output=True, timeout=30)
                
                if temp_pdf.exists() and HAS_PYPDF2:
                    text = ""
                    with open(temp_pdf, 'rb') as f:
                        reader = PyPDF2.PdfReader(f)
                        for page in reader.pages[:5]:
                            text += page.extract_text() + "\n"
                    
                    content = f"### {label} from {filepath.name} (CAJ converted)\n\n{text[:8000]}"
                    self.content.append(content)
                    print(f"   📄 Loaded CAJ (converted): {filepath.name}")
                    
                    # Clean up temp file
                    temp_pdf.unlink()
                    return
            except Exception as e:
                pass  # Fall through to metadata-only
        
        # If conversion fails, extract metadata only
        content = f"""### {label} from {filepath.name} (CAJ format)

**Note:** This is a CAJ (Chinese Academic Journal) file.

**File:** {filepath.name}
**Size:** {filepath.stat().st_size / 1024:.1f} KB

CAJ files require CNKI (China National Knowledge Infrastructure) tools to extract full text.
To use this reference, please:
1. Open in CAJViewer and export to PDF, or
2. Use caj2pdf tool: `pip install caj2pdf`

**Citation placeholder:** [{filepath.stem} - Chinese Academic Journal]
"""
        self.content.append(content)
        self.warnings.append(f"CAJ file requires manual conversion: {filepath.name}")
        print(f"   ⚠️  CAJ metadata only: {filepath.name}")
    
    def _load_text(self, filepath):
        """Load plain text file"""
        try:
            # Try different encodings
            encodings = ['utf-8', 'gbk', 'gb2312', 'latin-1']
            text = None
            
            for enc in encodings:
                try:
                    with open(filepath, 'r', encoding=enc) as f:
                        text = f.read()
                    break
                except UnicodeDecodeError:
                    continue
            
            if text is None:
                raise UnicodeDecodeError("All encodings failed")
            
            label = "Data" if self.is_data_dir else "Reference"
            content = f"### {label} from {filepath.name}\n\n{text[:8000]}"
            self.content.append(content)
            print(f"   📄 Loaded {label.lower()}: {filepath.name}")
        except Exception as e:
            self.warnings.append(f"Text error {filepath.name}: {e}")


def load_directory(data_dir=None, refs_dir=None):
    """Load data and references from directories"""
    data = ""
    refs = ""
    
    if data_dir:
        print(f"📂 Loading data from: {data_dir}")
        loader = SmartLoader(data_dir, is_data_dir=True)
        data = loader.load_all()
    
    if refs_dir:
        print(f"📂 Loading references from: {refs_dir}")
        loader = SmartLoader(refs_dir, is_data_dir=False)
        refs = loader.load_all()
    
    return data, refs


if __name__ == "__main__":
    # Test
    if len(sys.argv) > 1:
        data_dir = sys.argv[1] if len(sys.argv) > 1 else None
        refs_dir = sys.argv[2] if len(sys.argv) > 2 else None
        
        data, refs = load_directory(data_dir, refs_dir)
        
        print("\n" + "="*50)
        print("DATA:")
        print(data[:500] if data else "(No data)")
        print("\n" + "="*50)
        print("REFERENCES:")
        print(refs[:500] if refs else "(No references)")
